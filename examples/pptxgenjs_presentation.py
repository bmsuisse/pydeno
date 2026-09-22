"""Build a real PowerPoint deck with the real pptxgenjs bundle inside pydeno.

This is the full, end-to-end version of the pattern sketched in
`vendored_npm_libraries.py`: the host evaluates the pinned, **unmodified**
460,889-byte `pptxgenjs` 4.0.1 UMD bundle from `vendor/pptxgenjs/` in a bare
V8 isolate, injects three browser-global polyfills separately, and drives the
library's real API to produce a six-slide deck with a table, three native
OOXML charts (including a combo chart on a secondary axis) and an embedded
PNG.

Then it actually checks the result, in three escalating steps:

  1. `zipfile` -- integrity plus the expected OOXML parts.
  2. `python-pptx` -- a real third-party consumer opens the file and reads
     back the expected slides, table, charts and picture.
  3. LibreOffice + `pdftoppm` -- renders every slide to PNG so you can *look*
     at it. This step is the one that has historically caught defects the
     first two missed: an axis running negative, a legend in the wrong
     colour, a stray frame around a plot area, a scheme name in a literal
     colour slot coming out as bright cyan. Validation proves the file is
     well-formed; only rendering proves it looks right.

Run it:

    uv sync --group examples
    uv run python examples/pptxgenjs_presentation.py

No network access is needed -- the bundle is vendored. Step 3 is skipped with
a clear message if LibreOffice or Poppler are not installed; steps 1 and 2
always run.

The permanent, hermetic regression test for this chain is
`tests/test_vendored_bundle_execution.py`.
"""

from __future__ import annotations

import asyncio
import base64
import pathlib
import shutil
import struct
import subprocess
import sys
import zipfile
import zlib
from io import BytesIO

from pydeno import Runtime

REPO_ROOT = pathlib.Path(__file__).resolve().parent.parent
VENDOR = REPO_ROOT / "vendor" / "pptxgenjs"
BUNDLE_PATH = VENDOR / "pptxgen.bundle.js"
POLYFILLS_PATH = VENDOR / "polyfills.js"
OUT_DIR = REPO_ROOT / "build" / "pptxgenjs-example"


def make_png(width: int, height: int) -> bytes:
    """A real PNG, built with the standard library only (no Pillow needed)."""
    rows = bytearray()
    for y in range(height):
        rows.append(0)  # filter type: none
        for x in range(width):
            if height // 3 < y < 2 * height // 3:
                rows += bytes((20, 30, 60))
            else:
                rows += bytes(
                    (
                        int(255 * x / max(width - 1, 1)),
                        int(180 * y / max(height - 1, 1)),
                        140,
                    )
                )

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(bytes(rows), 9))
        + chunk(b"IEND", b"")
    )


# --- the deck --------------------------------------------------------------
#
# Two pptxgenjs footguns are deliberately avoided below. Both produce a file
# that looks fine to a structural check and is broken for real consumers:
#
#   * pptxgenjs MUTATES the options object you hand it, normalising units and
#     rewriting colour keys in place. Every add* call therefore gets a FRESH
#     object literal -- a shared `const opts = {...}` silently leaks state
#     from one shape into the next.
#   * `secondaryValAxis`/`secondaryCatAxis` need BOTH `valAxes` and `catAxes`
#     with two entries each. Supply only `valAxes` and PowerPoint discards
#     the chart and reports the whole file as corrupt.
#
# Three renderer-level lessons are also baked in, each marked at its line:
# always set both axis bounds, never put a theme-scheme name in a literal
# colour slot, and set legend/label colours explicitly.

DECK_SCRIPT_TEMPLATE = """
const pres = new PptxGenJS();
pres.layout = "LAYOUT_WIDE";
pres.author = "pydeno";
pres.title = "Real pptxgenjs inside pydeno";

// --- 1. title ---
const s1 = pres.addSlide();
s1.background = { color: "1F2A44" };
s1.addText("Real pptxgenjs, real V8, no Node", {
  x: 0.6, y: 2.0, w: 12.1, h: 1.0,
  fontSize: 34, bold: true, color: "FFFFFF", align: "center",
});
s1.addText("460,889 bytes of unmodified UMD bundle, three host polyfills", {
  x: 0.6, y: 3.1, w: 12.1, h: 0.6,
  fontSize: 16, color: "AAB6D3", align: "center",
});
s1.addNotes("Built inside a bare V8 isolate via pydeno, with no Node.js.");

// --- 2. table ---
const s2 = pres.addSlide();
s2.addText("Revenue by region", {
  x: 0.5, y: 0.35, w: 8, h: 0.6, fontSize: 24, bold: true, color: "1F2A44",
});
s2.addTable(
  [
    [
      { text: "Region", options: { bold: true, color: "FFFFFF", fill: { color: "1F2A44" } } },
      { text: "FY24", options: { bold: true, color: "FFFFFF", fill: { color: "1F2A44" } } },
      { text: "FY25", options: { bold: true, color: "FFFFFF", fill: { color: "1F2A44" } } },
    ],
    ["EMEA", "1.20M", "1.44M"],
    ["APAC", "0.90M", "1.17M"],
    ["AMER", "1.60M", "1.76M"],
  ],
  {
    x: 0.5, y: 1.2, w: 7.5, colW: [2.5, 2.5, 2.5],
    border: { pt: 1, color: "C8D0E0" }, fontSize: 14,
  }
);

// --- 3. bar chart ---
const s3 = pres.addSlide();
s3.addText("Bar chart (native OOXML chart part)", {
  x: 0.5, y: 0.35, w: 10, h: 0.6, fontSize: 24, bold: true, color: "1F2A44",
});
s3.addChart(
  pres.ChartType.bar,
  [{ name: "Revenue", labels: ["Q1", "Q2", "Q3", "Q4"], values: [10, 14, 13, 18] }],
  {
    x: 0.5, y: 1.2, w: 12.0, h: 5.0,
    showTitle: true, title: "Quarterly revenue", titleColor: "1F2A44", titleFontSize: 16,
    // Set legend/label colours explicitly: left unset, the legend inherits a
    // colour from whatever ran before it and can come out near-invisible.
    showLegend: true, legendPos: "b", legendColor: "363636",
    chartColors: ["4472C4"],
    catAxisLabelColor: "363636", valAxisLabelColor: "363636",
    // BOTH bounds. valAxisMaxVal alone lets the renderer pick its own
    // minimum, and an all-positive series ends up on a negative axis.
    valAxisMinVal: 0, valAxisMaxVal: 20,
    // Literal hex only. A scheme name here emits <a:srgbClr val="accent1"/>,
    // which is invalid OOXML and which LibreOffice draws as bright cyan.
    valGridLine: { style: "solid", size: 1, color: "E6E9F0" },
  }
);

// --- 4. line chart, two series ---
const s4 = pres.addSlide();
s4.addText("Line chart (two series)", {
  x: 0.5, y: 0.35, w: 10, h: 0.6, fontSize: 24, bold: true, color: "1F2A44",
});
s4.addChart(
  pres.ChartType.line,
  [
    { name: "Sessions", labels: ["Jan", "Feb", "Mar", "Apr", "May"], values: [4, 6, 9, 11, 15] },
    { name: "Signups", labels: ["Jan", "Feb", "Mar", "Apr", "May"], values: [1, 2, 4, 5, 8] },
  ],
  {
    x: 0.5, y: 1.2, w: 12.0, h: 5.0,
    showTitle: true, title: "Funnel trend", titleColor: "1F2A44", titleFontSize: 16,
    showLegend: true, legendPos: "b", legendColor: "363636",
    chartColors: ["4472C4", "ED7D31"],
    lineDataSymbol: "circle", lineSize: 2,
    catAxisLabelColor: "363636", valAxisLabelColor: "363636",
    valAxisMinVal: 0, valAxisMaxVal: 16,
    valGridLine: { style: "solid", size: 1, color: "E6E9F0" },
  }
);

// --- 5. combo chart on a secondary axis ---
const s5 = pres.addSlide();
s5.addText("Combo chart (secondary value axis)", {
  x: 0.5, y: 0.35, w: 10, h: 0.6, fontSize: 24, bold: true, color: "1F2A44",
});
s5.addChart(
  [
    {
      type: pres.ChartType.bar,
      data: [{ name: "Units", labels: ["Q1", "Q2", "Q3", "Q4"], values: [120, 150, 140, 190] }],
      options: { chartColors: ["4472C4"], barGrouping: "clustered" },
    },
    {
      type: pres.ChartType.line,
      data: [{ name: "Margin %", labels: ["Q1", "Q2", "Q3", "Q4"], values: [12, 15, 14, 21] }],
      options: {
        chartColors: ["ED7D31"],
        secondaryValAxis: true,
        secondaryCatAxis: true,
        lineDataSymbol: "circle",
        lineSize: 2,
      },
    },
  ],
  {
    x: 0.5, y: 1.2, w: 12.0, h: 5.0,
    showTitle: true, title: "Units vs margin", titleColor: "1F2A44", titleFontSize: 16,
    showLegend: true, legendPos: "b", legendColor: "363636",
    // Two valAxes AND two catAxes -- see the note above.
    valAxes: [
      {
        showValAxisTitle: true, valAxisTitle: "Units",
        valAxisLabelColor: "363636", valAxisTitleColor: "363636",
        valAxisMinVal: 0, valAxisMaxVal: 200,
        valGridLine: { style: "solid", size: 1, color: "E6E9F0" },
      },
      {
        showValAxisTitle: true, valAxisTitle: "Margin %",
        valAxisLabelColor: "363636", valAxisTitleColor: "363636",
        valAxisMinVal: 0, valAxisMaxVal: 25,
        // No second set of grid lines -- they would cross the first set.
        valGridLine: { style: "none" },
      },
    ],
    catAxes: [
      { catAxisTitle: "Quarter", catAxisLabelColor: "363636" },
      { catAxisHidden: true },
    ],
  }
);

// --- 6. embedded image ---
const s6 = pres.addSlide();
s6.addText("Embedded image (base64 into ppt/media)", {
  x: 0.5, y: 0.35, w: 10, h: 0.6, fontSize: 24, bold: true, color: "1F2A44",
});
s6.addImage({
  data: "image/png;base64,__IMAGE_B64__",
  x: 0.5, y: 1.3, w: 6.0, h: 4.0,
});
s6.addText("A real PNG part, carried through the bundle byte-for-byte.", {
  x: 7.0, y: 1.3, w: 5.6, h: 2.0, fontSize: 15, color: "363636",
});

pres.write({ outputType: "base64" });
"""


async def build_deck() -> bytes:
    """Evaluate the polyfills, then the bundle, then drive the real API."""
    polyfills = POLYFILLS_PATH.read_text()
    bundle = BUNDLE_PATH.read_text()
    script = DECK_SCRIPT_TEMPLATE.replace(
        "__IMAGE_B64__", base64.b64encode(make_png(360, 240)).decode()
    )

    print(
        f"  bundle: {BUNDLE_PATH.relative_to(REPO_ROOT)} "
        f"({len(bundle.encode()):,} bytes, unmodified)"
    )

    with Runtime() as runtime:
        # The polyfills go in SEPARATELY. The bundle is never patched.
        runtime.eval(polyfills)
        runtime.eval(bundle)

        # The UMD preamble has to have installed these for anything to work.
        assert runtime.eval("typeof JSZip") == "function"
        assert runtime.eval("typeof PptxGenJS") == "function"
        print("  bundle evaluated; JSZip + PptxGenJS globals installed")

        # write() is async: JSZip advances its pipeline through setTimeout,
        # so this only ever resolves because the polyfill defers to the
        # microtask queue.
        b64 = await runtime.eval_async(script, timeout=120.0)

    return base64.b64decode(b64)


def validate_zip(raw: bytes) -> None:
    """Step 1: the standard library's view of the file."""
    zf = zipfile.ZipFile(BytesIO(raw))
    if zf.testzip() is not None:
        raise AssertionError("corrupt member in the .pptx zip")

    names = zf.namelist()
    slides = [
        n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")
    ]
    charts = [
        n for n in names if n.startswith("ppt/charts/chart") and n.endswith(".xml")
    ]
    media = [n for n in names if n.startswith("ppt/media/") and not n.endswith("/")]

    assert "[Content_Types].xml" in names, "not an OOXML package"
    assert len(slides) == 6, f"expected 6 slides, got {len(slides)}"
    assert len(charts) == 3, f"expected 3 chart parts, got {len(charts)}"
    assert len(media) == 1, f"expected 1 media part, got {len(media)}"
    assert zf.read(media[0]).startswith(b"\x89PNG"), "media part is not a PNG"

    # The combo chart must have two value axes AND two category axes.
    combo = [n for n in charts if zf.read(n).decode().count("<c:valAx>") == 2]
    assert len(combo) == 1, "combo chart with a secondary value axis not found"
    xml = zf.read(combo[0]).decode()
    assert xml.count("<c:catAx>") == 2, (
        f"{combo[0]}: 2 <c:valAx> but {xml.count('<c:catAx>')} <c:catAx> -- "
        "PowerPoint would report this file as corrupt"
    )

    print(
        f"  zipfile: {len(names)} parts, {len(slides)} slides, "
        f"{len(charts)} charts, {len(media)} image, combo chart axes matched"
    )


def validate_python_pptx(path: pathlib.Path) -> None:
    """Step 2: a real third-party consumer reads the file back."""
    try:
        from pptx import Presentation
    except ImportError:
        print("  python-pptx: NOT INSTALLED -- run `uv sync --group examples`")
        return

    prs = Presentation(str(path))
    tables = charts = pictures = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                tables += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if "PICTURE" in str(shape.shape_type):
                pictures += 1

    assert len(prs.slides) == 6, f"python-pptx sees {len(prs.slides)} slides"
    assert tables == 1, f"tables: {tables}"
    assert charts == 3, f"charts: {charts}"
    assert pictures == 1, f"pictures: {pictures}"
    print(
        f"  python-pptx: 6 slides, {tables} table, {charts} charts, "
        f"{pictures} picture -- opens cleanly"
    )


def render_to_png(path: pathlib.Path) -> list[pathlib.Path]:
    """Step 3: render every slide so a human can look at it."""
    soffice = shutil.which("soffice") or shutil.which(
        "soffice", path="/Applications/LibreOffice.app/Contents/MacOS"
    )
    pdftoppm = shutil.which("pdftoppm")

    if not soffice or not pdftoppm:
        missing = []
        if not soffice:
            missing.append("LibreOffice (`soffice`)")
        if not pdftoppm:
            missing.append("Poppler (`pdftoppm`)")
        print(f"  render: SKIPPED -- {' and '.join(missing)} not installed.")
        print(
            "          Validation above still ran; only the visual check was skipped."
        )
        return []

    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(path.parent),
            str(path),
        ],
        check=True,
        capture_output=True,
        timeout=300,
    )
    pdf = path.with_suffix(".pdf")
    subprocess.run(
        [pdftoppm, "-png", "-r", "150", str(pdf), str(path.parent / "slide")],
        check=True,
        capture_output=True,
        timeout=300,
    )
    pngs = sorted(path.parent.glob("slide-*.png"))
    print(f"  render: {pdf.name} -> {len(pngs)} PNGs at 150 dpi")
    return pngs


async def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    print("Building a six-slide deck with the vendored pptxgenjs 4.0.1 bundle...")
    raw = await build_deck()

    path = OUT_DIR / "pydeno_pptxgenjs.pptx"
    path.write_bytes(raw)
    print(f"\nWrote {path} ({len(raw):,} bytes)\n")

    print("Validating:")
    validate_zip(raw)
    validate_python_pptx(path)
    pngs = render_to_png(path)

    print(f"\nDeck:  {path}")
    for png in pngs:
        print(f"Slide: {png}")
    if pngs:
        print(
            "\nOpen the PNGs and actually look at them -- rendering catches "
            "defects that\nstructural validation cannot: axes running "
            "negative, legends in the wrong\ncolour, stray frames, invalid "
            "colours drawn as bright cyan."
        )
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
